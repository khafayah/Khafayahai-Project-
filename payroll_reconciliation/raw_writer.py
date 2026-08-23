"""Writes the six raw iTrent paste tabs into worksheets, styled per Hakim
Group branding. Used to populate Output 1's paste tabs with a working
demonstration dataset (see mock_data.py) so the delivered tool is a live
example, not an empty shell.
"""
from openpyxl.utils import column_index_from_string, get_column_letter
import styles as S
import raw_layout as L
from mock_data import ALL_ELEMENTS


def _sheet_intro(ws, title, note_row, note_text):
    ws.sheet_view.showGridLines = False
    S.apply_offwhite_background(ws, max_row=250, max_col=45)
    ws["A1"] = title
    ws["A1"].font = S.f_subheader()
    ws["A" + str(note_row)] = note_text
    ws["A" + str(note_row)].font = S.f_note()


def _write_header(ws, header_row, columns: dict):
    for letter, name in columns.items():
        cell = ws[f"{letter}{header_row}"]
        cell.value = name
    last_col_idx = max(column_index_from_string(l) for l in columns)
    S.style_header_row(ws, header_row, 1, last_col_idx)
    return last_col_idx


def write_npv_comparison(ws, dataset):
    _sheet_intro(ws, "NPV Comparison (Net Pay) - paste raw iTrent export below from row 14",
                 12, "Paste the full export starting at A14 (headers) so data begins row 15. Do not insert or delete columns.")
    last_col = _write_header(ws, L.NPV_HEADER_ROW, L.NPV_COLUMNS)
    row = L.NPV_DATA_START
    for emp in dataset["employees"]:
        surname = emp.payroll_number  # synthetic - real export carries the employee's name
        ws[f"A{row}"] = emp.practice[:3].upper()
        ws[f"B{row}"] = f"Surname{emp.payroll_number[-3:]}"
        ws[f"C{row}"] = f"Forename{emp.payroll_number[-3:]}"
        ws[f"D{row}"] = emp.payroll_number
        ws[f"E{row}"] = emp.prior_net if emp.prior_net is not None else None
        ws[f"F{row}"] = emp.current_net if emp.current_net is not None else None
        ws[f"G{row}"] = f"=IF(AND(E{row}<>\"\",F{row}<>\"\"),F{row}-E{row},IF(F{row}<>\"\",F{row},IF(E{row}<>\"\",-E{row},\"\")))"
        ws[f"H{row}"] = f"=IFERROR(G{row}/E{row},\"\")"
        ws[f"I{row}"] = dataset["period_current"]
        ws[f"J{row}"] = "Monthly"
        for c in range(1, last_col + 1):
            cell = ws.cell(row=row, column=c)
            fmt = None
            if c in (5, 6, 7):
                fmt = '#,##0.00;(#,##0.00)'
            elif c == 8:
                fmt = '0.0%'
            S.style_body_cell(cell, number_format=fmt)
        row += 1
    last_row = row - 1
    S.col_widths(ws, {"A": 10, "B": 14, "C": 14, "D": 14, "E": 13, "F": 13, "G": 12, "H": 11, "I": 13, "J": 11})
    S.freeze_and_filter(ws, L.NPV_HEADER_ROW, 1, last_col, last_row)
    S.set_tab_color(ws, secondary=True)
    return last_row


def write_payment_breakdown(ws, dataset, period):
    label = "current" if period == "current" else "prior"
    warn = ""
    if period == "prior":
        warn = ("If this export carries an EXTRA leading 'Reference No' column shifting everything "
                "right by one, delete that column before pasting - see Control Sheet warning.")
    _sheet_intro(ws, f"Payment Breakdown - {label} month ({dataset['period_' + period]}) - paste from row 14",
                 12, warn or "Paste the full export starting at A14 (headers) so data begins row 15.")
    columns = dict(L.PB_PREFIX_COLUMNS)
    for i, name in enumerate(ALL_ELEMENTS):
        columns[L.pb_element_col_letter(i)] = name
    columns[L.PB_REFNO_LETTER] = "Reference No"
    last_col = _write_header(ws, L.PB_HEADER_ROW, columns)

    row = L.PB_DATA_START
    for emp in dataset["employees"]:
        net = emp.current_net if period == "current" else emp.prior_net
        if net is None:
            continue
        ws[f"A{row}"] = row - L.PB_DATA_START + 1
        ws[f"B{row}"] = f"Surname{emp.payroll_number[-3:]}"
        ws[f"C{row}"] = f"Forename{emp.payroll_number[-3:]}"
        ws[f"D{row}"] = emp.practice
        ws[f"E{row}"] = "Clinical" if emp.category == "Practice Clinical" else "Support"
        ws[f"F{row}"] = "CC-" + emp.practice[:3].upper()
        ws[f"G{row}"] = dataset["period_current"] if period == "current" else dataset["period_prior"]
        ws[f"H{row}"] = "Monthly"
        ws[f"I{row}"] = "1257L"
        ws[f"J{row}"] = "A"
        idx = 0 if period == "prior" else 1
        for i, name in enumerate(ALL_ELEMENTS):
            amt = emp.elements.get(name, (0.0, 0.0))[idx]
            letter = L.pb_element_col_letter(i)
            ws[f"{letter}{row}"] = amt if amt else 0
        ws[f"{L.PB_REFNO_LETTER}{row}"] = emp.payroll_number
        for c in range(1, last_col + 1):
            cell = ws.cell(row=row, column=c)
            fmt = '#,##0.00' if c >= L.PB_ELEMENT_START_COL and c <= L.PB_ELEMENT_END_COL else None
            S.style_body_cell(cell, number_format=fmt)
        row += 1
    last_row = row - 1
    widths = {"A": 6, "B": 14, "C": 14, "D": 22, "E": 10, "F": 10, "G": 12, "H": 10, "I": 9, "J": 9}
    for i in range(len(ALL_ELEMENTS)):
        widths[L.pb_element_col_letter(i)] = 13
    widths[L.PB_REFNO_LETTER] = 14
    S.col_widths(ws, widths)
    S.freeze_and_filter(ws, L.PB_HEADER_ROW, 1, last_col, max(last_row, L.PB_DATA_START))
    S.set_tab_color(ws, secondary=True)
    return last_row


def write_new_starters(ws, dataset):
    _sheet_intro(ws, "New Starters - paste raw iTrent export below from row 13",
                 11, "Paste the full export starting at A13 (headers) so data begins row 14.")
    last_col = _write_header(ws, L.NS_HEADER_ROW, L.NS_COLUMNS)
    row = L.NS_DATA_START
    for emp in dataset["starters"]:
        ws[f"A{row}"] = row - L.NS_DATA_START + 1
        ws[f"B{row}"] = "Mr/Mrs"
        ws[f"C{row}"] = f"Surname{emp.payroll_number[-3:]}"
        ws[f"D{row}"] = emp.payroll_number
        ws[f"E{row}"] = emp.practice
        ws[f"F{row}"] = emp.position_name
        ws[f"G{row}"] = dataset["period_current"]
        ws[f"H{row}"] = emp.category
        ws[f"I{row}"] = "/".join(emp.basis_list)
        ws[f"J{row}"] = "N/A"
        ws[f"K{row}"] = emp.personal_reference
        for c in range(1, last_col + 1):
            S.style_body_cell(ws.cell(row=row, column=c))
        row += 1
    last_row = max(row - 1, L.NS_DATA_START)
    S.col_widths(ws, {"A": 6, "B": 8, "C": 14, "D": 14, "E": 22, "F": 20, "G": 12, "H": 16, "I": 12, "J": 12, "K": 14})
    S.freeze_and_filter(ws, L.NS_HEADER_ROW, 1, last_col, last_row)
    S.set_tab_color(ws, secondary=True)
    return last_row


def write_leavers(ws, dataset):
    _sheet_intro(ws, "Leavers - paste raw iTrent export below from row 11",
                 9, "Paste the full export starting at A11 (headers) so data begins row 12.")
    last_col = _write_header(ws, L.LV_HEADER_ROW, L.LV_COLUMNS)
    row = L.LV_DATA_START
    for emp in dataset["leavers"]:
        ws[f"A{row}"] = row - L.LV_DATA_START + 1
        ws[f"B{row}"] = "Mr/Mrs"
        ws[f"C{row}"] = f"Surname{emp.payroll_number[-3:]}"
        ws[f"D{row}"] = emp.payroll_number
        ws[f"E{row}"] = emp.practice
        ws[f"F{row}"] = emp.position_name
        ws[f"G{row}"] = dataset["period_current"]
        ws[f"H{row}"] = emp.category
        ws[f"I{row}"] = "/".join(emp.basis_list)
        ws[f"J{row}"] = "N/A"
        ws[f"K{row}"] = emp.personal_reference
        for c in range(1, last_col + 1):
            S.style_body_cell(ws.cell(row=row, column=c))
        row += 1
    last_row = max(row - 1, L.LV_DATA_START)
    S.col_widths(ws, {"A": 6, "B": 8, "C": 14, "D": 14, "E": 22, "F": 20, "G": 12, "H": 16, "I": 12, "J": 12, "K": 14})
    S.freeze_and_filter(ws, L.LV_HEADER_ROW, 1, last_col, last_row)
    S.set_tab_color(ws, secondary=True)
    return last_row


def write_structure(ws, dataset):
    _sheet_intro(ws, "Structure Report - paste raw iTrent export below from row 12",
                 10, "Paste the full export starting at A12 (headers) so data begins row 13. One row per assignment - staff with more than one contract (e.g. Zero Hours + Permanent) appear on multiple rows.")
    columns = dict(L.ST_PREFIX_COLUMNS)
    from mock_data import STRUCTURE_FILLER_COLS
    for i, name in enumerate(STRUCTURE_FILLER_COLS):
        columns[get_column_letter(L.ST_FILLER_START_COL + i)] = name
    columns[L.ST_CATEGORY_COL] = "Category"
    last_col = _write_header(ws, L.ST_HEADER_ROW, columns)

    row = L.ST_DATA_START
    for emp in dataset["employees"]:
        for basis in emp.basis_list:
            first, last = emp.payroll_number, emp.payroll_number  # synthetic
            ws[f"A{row}"] = row - L.ST_DATA_START + 1
            ws[f"B{row}"] = "Mr/Mrs"
            ws[f"C{row}"] = f"Forename{emp.payroll_number[-3:]}"
            ws[f"D{row}"] = f"Surname{emp.payroll_number[-3:]}"
            ws[f"E{row}"] = emp.practice
            ws[f"F{row}"] = "Clinical" if emp.category == "Practice Clinical" else "Support"
            ws[f"G{row}"] = "CC-" + emp.practice[:3].upper()
            ws[f"H{row}"] = "01/01/2022"
            ws[f"I{row}"] = "" if not emp.is_leaver else dataset["period_current"]
            ws[f"J{row}"] = "Leaver" if emp.is_leaver else "Active"
            ws[f"K{row}"] = 37.5 if basis != "Zero Hours" else 0
            ws[f"L{row}"] = emp.position_name
            ws[f"M{row}"] = "Band 4"
            ws[f"N{row}"] = basis
            ws[f"O{row}"] = 1.0 if basis != "Zero Hours" else 0.0
            ws[f"P{row}"] = emp.payroll_number
            for i, val in enumerate(["N/A"] * len(STRUCTURE_FILLER_COLS)):
                ws[f"{get_column_letter(L.ST_FILLER_START_COL + i)}{row}"] = val
            ws[f"{L.ST_CATEGORY_COL}{row}"] = emp.category
            for c in range(1, last_col + 1):
                S.style_body_cell(ws.cell(row=row, column=c))
            row += 1
    last_row = row - 1
    widths = {"A": 6, "B": 8, "C": 14, "D": 14, "E": 22, "F": 10, "G": 10, "H": 11, "I": 11, "J": 10,
              "K": 10, "L": 22, "M": 9, "N": 13, "O": 7, "P": 14}
    S.col_widths(ws, widths)
    S.freeze_and_filter(ws, L.ST_HEADER_ROW, 1, last_col, last_row)
    S.set_tab_color(ws, secondary=True)
    return last_row
