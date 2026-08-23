"""End-to-end build + validation for the Hakim Group payroll reconciliation pack.

Builds all three workbooks from the (synthetic, demonstration) source data,
recalculates each with LibreOffice, asserts zero formula errors, and
cross-checks the workbooks' own computed KPIs against reference_calc.py's
independent Python calculation. Run this after any change to output1/2/3.py.
"""
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from mock_data import build_dataset
from reference_calc import compute_reconciliation
from output1 import build_output1
from output2 import build_output2
from output3 import build_output3

RECALC_SCRIPT = "/root/.claude/skills/synced/xlsx/scripts/recalc.py"
OUT_DIR = Path(__file__).parent / "output"

FILES = {
    "output1": OUT_DIR / "Hakim Group - Payroll Reconciliation Tool.xlsx",
    "output2": OUT_DIR / "Hakim Group - Investigation List.xlsx",
    "output3": OUT_DIR / "Hakim Group - SLT Pack.xlsx",
}


def recalc(path, timeout=60):
    result = subprocess.run(
        [sys.executable, RECALC_SCRIPT, str(path), str(timeout)],
        capture_output=True, text=True,
    )
    import json
    return json.loads(result.stdout)


def check(label, actual, expected):
    ok = actual == expected
    mark = "OK " if ok else "FAIL"
    print(f"  [{mark}] {label}: workbook={actual!r} reference={expected!r}")
    return ok


def main():
    OUT_DIR.mkdir(exist_ok=True)
    dataset = build_dataset()
    result = compute_reconciliation(dataset)
    s = result["summary"]

    print("Building workbooks...")
    build_output1(dataset, str(FILES["output1"]))
    build_output2(result, str(FILES["output2"]))
    build_output3(dataset, result, str(FILES["output3"]))

    all_ok = True
    for key, path in FILES.items():
        print(f"\nRecalculating {path.name} ...")
        r = recalc(path)
        if r.get("status") != "success" or r.get("total_errors", 1) != 0:
            print(f"  FAIL: {r}")
            all_ok = False
        else:
            print(f"  OK: {r['total_formulas']} formulas, 0 errors")

    print("\nCross-checking Output 1 (Reconciliation Tool) Summary KPIs...")
    wb1 = load_workbook(FILES["output1"], data_only=True)
    sm = wb1["Summary"]
    checks = [
        ("Employee count", sm["D5"].value, s["employee_count"]),
        ("Movements over tolerance", sm["D6"].value, s["movements_over_tolerance"]),
        ("Explained starters/leavers", sm["D7"].value, s["explained_starters_leavers"]),
        ("Zero hours count", sm["D8"].value, s["zero_hours_count"]),
        ("Review required count", sm["D9"].value, s["review_count"]),
    ]
    for label, actual, expected in checks:
        all_ok &= check(label, actual, expected)

    print("\nCross-checking Output 2 (Investigation List) Summary KPIs...")
    wb2 = load_workbook(FILES["output2"], data_only=True)
    sm2 = wb2["Summary"]
    priority_counts = {"High": 0, "Medium": 0, "Low": 0}
    for item in result["review_items"]:
        priority_counts[item["priority"]] += 1
    checks2 = [
        ("Total review items", sm2["D5"].value, s["review_count"]),
        ("High priority", sm2["D6"].value, priority_counts["High"]),
        ("Medium priority", sm2["D7"].value, priority_counts["Medium"]),
        ("Low priority", sm2["D8"].value, priority_counts["Low"]),
    ]
    for label, actual, expected in checks2:
        all_ok &= check(label, actual, expected)

    print("\nCross-checking Output 3 (SLT Pack) Dashboard KPIs...")
    wb3 = load_workbook(FILES["output3"], data_only=True)
    d3 = wb3["Dashboard"]
    checks3 = [
        ("Employees compared", d3["A5"].value, s["employee_count"]),
        ("Over tolerance", d3["I5"].value, s["movements_over_tolerance"]),
        ("Review required", d3["K5"].value, s["review_count"]),
    ]
    for label, actual, expected in checks3:
        all_ok &= check(label, actual, expected)

    print("\n" + ("ALL CHECKS PASSED" if all_ok else "SOME CHECKS FAILED"))
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()
