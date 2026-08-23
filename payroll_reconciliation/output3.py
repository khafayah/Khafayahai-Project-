"""Builds Output 3: the SLT Pack - GDPR-safe (payroll numbers only, no names)."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import DoughnutChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
import styles as S

SORT_RANK = {"Review required": 0, "Explained: Starter": 1, "Explained: Leaver": 1,
             "Explained: Paired transfer": 1, "OK": 2, "Zero Hours": 3}

REASON_TEXT = {
    "Explained: Starter": "New starter this period",
    "Explained: Leaver": "Leaver this period",
    "Explained: Paired transfer": "Pay moved to a different payroll record for this person",
    "OK": "Movement within tolerance - no action needed",
    "Zero Hours": "Zero hours contract - variable pay, not reviewed",
}


def _all_employee_rows(result):
    rows = []
    for r in result["main_rows"]:
        reason = REASON_TEXT.get(r["status"])
        if r["status"] == "Review required":
            cause = r["cause"]["element"] if r["cause"] else None
            reason = f"Review required - main driver: {cause}" if cause else "Review required"
        rows.append({
            "payroll_number": r["payroll_number"], "practice": r["practice"],
            "prior_net": r["prior_net"], "current_net": r["current_net"],
            "diff": r["diff"], "diff_pct": r["diff_pct"], "reason": reason,
            "status": r["status"],
        })
    for z in result["zero_hours_rows"]:
        rows.append({
            "payroll_number": z["payroll_number"], "practice": z["practice"],
            "prior_net": z["prior_net"], "current_net": z["current_net"],
            "diff": z["diff"], "diff_pct": z["diff_pct"],
            "reason": REASON_TEXT["Zero Hours"], "status": "Zero Hours",
        })
    rows.sort(key=lambda r: (SORT_RANK.get(r["status"], 2), -abs(r["diff"] or 0)))
    return rows


def build_all_employees_sheet(wb, result):
    ws = wb.create_sheet("All Employees")
    ws.sheet_view.showGridLines = False
    rows = _all_employee_rows(result)
    S.apply_offwhite_background(ws, max_row=6 + len(rows), max_col=8)
    S.set_tab_color(ws)
    ws.merge_cells("A1:G1")
    S.style_title(ws["A1"], "All Employees - Payroll Numbers Only")
    ws["A2"] = "GDPR: no employee names are held in this pack. Sorted review-first."
    ws["A2"].font = S.f_note()
    ws.merge_cells("A2:G2")

    header_row = 3
    headers = ["Payroll Number", "Practice", "Prior Net", "Current Net", "Difference £",
               "Difference %", "Reason"]
    for i, h in enumerate(headers):
        ws.cell(row=header_row, column=1 + i, value=h)
    S.style_header_row(ws, header_row, 1, len(headers))
    data_start = header_row + 1
    for i, r in enumerate(rows):
        rr = data_start + i
        vals = [r["payroll_number"], r["practice"], r["prior_net"], r["current_net"],
                r["diff"], r["diff_pct"], r["reason"]]
        for j, v in enumerate(vals):
            cell = ws.cell(row=rr, column=1 + j, value=v)
            cell.font = S.f_body()
            cell.border = S.BORDER_THIN
            cell.alignment = Alignment(horizontal="left" if j in (0, 1, 6) else "center")
        for col_idx, fmt in [(3, '#,##0.00'), (4, '#,##0.00'), (5, '#,##0.00'), (6, '0.0%')]:
            ws.cell(row=rr, column=col_idx).number_format = fmt
    last_row = data_start + len(rows) - 1 if rows else data_start
    S.add_status_conditional_formatting(ws, "G", data_start, last_row, "A", "G")
    ws.conditional_formatting.add(
        f"G{data_start}:G{last_row}",
        S.FormulaRule(formula=[f'LEFT($G{data_start},16)="Review required"'], fill=S.FILL_PINK_LIGHT,
                      font=Font(name=S.FONT_NAME, color="C00000", bold=True)))
    S.col_widths(ws, {"A": 14, "B": 24, "C": 12, "D": 12, "E": 12, "F": 11, "G": 44})
    S.freeze_and_filter(ws, header_row, 1, len(headers), last_row)
    return last_row, data_start


def build_dashboard_sheet(wb, result, all_emp_last_row, all_emp_data_start):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    S.apply_offwhite_background(ws, max_row=55, max_col=16)
    S.set_tab_color(ws)
    ws.merge_cells("A1:P2")
    S.style_title(ws["A1"], "Hakim Group  |  Payroll SLT Pack")

    s = result["summary"]
    cards = [
        ("Employees compared", f"={s['employee_count']}", "#,##0"),
        ("Prior net pay", f"={s['total_net_prior']}", '£#,##0'),
        ("Current net pay", f"={s['total_net_current']}", '£#,##0'),
        ("Net movement", f"={round(s['total_net_current'] - s['total_net_prior'], 2)}", '+£#,##0;-£#,##0'),
        ("Over tolerance", f"={s['movements_over_tolerance']}", "#,##0"),
        ("Review required", f"={s['review_count']}", "#,##0"),
    ]
    col_width = 2
    start_col = 1
    for i, (label, formula, fmt) in enumerate(cards):
        c = start_col + i * col_width
        ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + col_width - 1)
        ws.merge_cells(start_row=5, start_column=c, end_row=7, end_column=c + col_width - 1)
        label_cell = ws.cell(row=4, column=c)
        label_cell.value = label
        label_cell.font = S.f_kpi_label()
        label_cell.fill = S.FILL_MID_PURPLE
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        val_cell = ws.cell(row=5, column=c)
        val_cell.value = formula
        val_cell.number_format = fmt
        val_cell.font = S.f_kpi_value()
        val_cell.fill = S.FILL_PURPLE
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 30
    ws.row_dimensions[5].height = 30

    # ---------------- Status donut ----------------
    donut_row = 10
    ws[f"A{donut_row}"] = "Status breakdown"
    ws[f"A{donut_row}"].font = S.f_subheader()
    status_header = donut_row + 1
    ws.cell(row=status_header, column=1, value="Status")
    ws.cell(row=status_header, column=2, value="Count")
    S.style_header_row(ws, status_header, 1, 2)
    status_counts = {
        "Review required": s["review_count"],
        "Explained (starter/leaver/transfer)": s["explained_starters_leavers"],
        "OK - within tolerance": s["employee_count"] - s["review_count"] - s["explained_starters_leavers"] + 0,
        "Zero hours": s["zero_hours_count"],
    }
    for i, (k, v) in enumerate(status_counts.items()):
        r = status_header + 1 + i
        ws.cell(row=r, column=1, value=k).font = S.f_body()
        cell = ws.cell(row=r, column=2, value=v)
        cell.number_format = "#,##0"
        cell.border = S.BORDER_THIN
    status_last = status_header + len(status_counts)

    donut = DoughnutChart()
    donut.title = "Status breakdown"
    donut.height, donut.width = 8, 10
    data = Reference(ws, min_col=2, min_row=status_header, max_row=status_last)
    cats = Reference(ws, min_col=1, min_row=status_header + 1, max_row=status_last)
    donut.add_data(data, titles_from_data=True)
    donut.set_categories(cats)
    donut.dataLabels = DataLabelList()
    donut.dataLabels.showPercent = True
    ws.add_chart(donut, "D10")

    # ---------------- Cause bar chart ----------------
    cause_row = status_header
    cause_col_start = 9
    ws.cell(row=donut_row, column=cause_col_start, value="Review causes").font = S.f_subheader()
    cause_header = donut_row + 1
    ws.cell(row=cause_header, column=cause_col_start, value="Cause")
    ws.cell(row=cause_header, column=cause_col_start + 1, value="Count")
    S.style_header_row(ws, cause_header, cause_col_start, cause_col_start + 1)
    causes = {}
    for item in result["review_items"]:
        name = item["cause"]["element"] if item["cause"] else "Unmatched"
        causes[name] = causes.get(name, 0) + 1
    if not causes:
        causes = {"None": 0}
    for i, (k, v) in enumerate(sorted(causes.items(), key=lambda x: -x[1])):
        r = cause_header + 1 + i
        ws.cell(row=r, column=cause_col_start, value=k).font = S.f_body()
        cell = ws.cell(row=r, column=cause_col_start + 1, value=v)
        cell.number_format = "#,##0"
        cell.border = S.BORDER_THIN
    cause_last = cause_header + max(len(causes), 1)

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Review causes"
    bar.height, bar.width = 8, 10
    bdata = Reference(ws, min_col=cause_col_start + 1, min_row=cause_header, max_row=cause_last)
    bcats = Reference(ws, min_col=cause_col_start, min_row=cause_header + 1, max_row=cause_last)
    bar.add_data(bdata, titles_from_data=True)
    bar.set_categories(bcats)
    bar.legend = None
    ws.add_chart(bar, "L10")

    # ---------------- Top review movements table ----------------
    top_row = status_last + 3
    ws[f"A{top_row}"] = "Top review movements"
    ws[f"A{top_row}"].font = S.f_subheader()
    thdr = top_row + 1
    headers = ["Payroll Number", "Practice", "Difference £", "Priority", "Cause"]
    for i, h in enumerate(headers):
        ws.cell(row=thdr, column=1 + i, value=h)
    S.style_header_row(ws, thdr, 1, len(headers))
    top10 = result["top20"][:10]
    for i, item in enumerate(top10):
        r = thdr + 1 + i
        cause = item["cause"]["element"] if item["cause"] else "Unmatched"
        vals = [item["payroll_number"], item["practice"], item["diff"], item["priority"], cause]
        for j, v in enumerate(vals):
            cell = ws.cell(row=r, column=1 + j, value=v)
            cell.font = S.f_body()
            cell.border = S.BORDER_THIN
        ws.cell(row=r, column=3).number_format = '£#,##0.00'
    top_last = thdr + max(len(top10), 1)
    S.add_status_conditional_formatting(ws, "D", thdr + 1, top_last, "A", "E")

    # ---------------- Assurance line ----------------
    assure_row = top_last + 3
    ws.merge_cells(f"A{assure_row}:P{assure_row + 3}")
    assure_cell = ws.cell(row=assure_row, column=1)
    assure_cell.value = (
        f"Assurance: {s['employee_count']} payroll records compared. Total Gross moved from "
        f"£{s['total_gross_prior']:,.0f} to £{s['total_gross_current']:,.0f}. Employer NI "
        f"£{s['employer_ni_prior']:,.0f} -> £{s['employer_ni_current']:,.0f}, Employer Pension "
        f"£{s['employer_pension_prior']:,.0f} -> £{s['employer_pension_current']:,.0f}, "
        f"Apprenticeship Levy £{s['apprenticeship_levy_prior']:,.0f} -> £{s['apprenticeship_levy_current']:,.0f}. "
        f"Prior records {s['prior_records']}, current records {s['current_records']} "
        f"({'headcount reconciles' if s['headcount_check'] else 'headcount check required'})."
    )
    assure_cell.font = Font(name=S.FONT_NAME, size=10, italic=True, color="FFFFFF")
    assure_cell.fill = S.FILL_MID_PURPLE
    assure_cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1)

    ws.freeze_panes = "A3"
    S.col_widths(ws, {c: 12 for c in "ABCDEFGHIJKLMNOP"})
    return ws


def build_expenses_sheet(wb, dataset, result):
    ws = wb.create_sheet("Expenses")
    ws.sheet_view.showGridLines = False
    S.apply_offwhite_background(ws, max_row=55, max_col=12)
    S.set_tab_color(ws)
    ws.merge_cells("A1:J1")
    S.style_title(ws["A1"], "Expenses - Current Month")

    employees = dataset["employees"]
    has_mileage = any(e.elements.get("Mileage Expenses", (0, 0))[1] for e in employees)
    note = "Mileage Expenses element found for this payroll and is included below." if has_mileage else \
        "Mileage Expenses element not present for this payroll - excluded."
    ws["A2"] = note
    ws["A2"].font = S.f_note()
    ws.merge_cells("A2:J2")

    def expense_total(e):
        total = e.elements.get("Other Expenses", (0, 0))[1]
        if has_mileage:
            total += e.elements.get("Mileage Expenses", (0, 0))[1]
        return round(total, 2)

    by_practice = {}
    by_category = {}
    by_position = {}
    for e in employees:
        amt = expense_total(e)
        if amt <= 0:
            continue
        by_practice[e.practice] = round(by_practice.get(e.practice, 0) + amt, 2)
        by_category[e.category] = round(by_category.get(e.category, 0) + amt, 2)
        cur_total, cur_count = by_position.get(e.position_name, (0, 0))
        by_position[e.position_name] = (round(cur_total + amt, 2), cur_count + 1)

    # ---------------- By practice: top 20 + rest grouped ----------------
    prac_row = 4
    ws[f"A{prac_row}"] = "By practice"
    ws[f"A{prac_row}"].font = S.f_subheader()
    phdr = prac_row + 1
    ws.cell(row=phdr, column=1, value="Practice")
    ws.cell(row=phdr, column=2, value="Total Expenses £")
    S.style_header_row(ws, phdr, 1, 2)
    ranked = sorted(by_practice.items(), key=lambda x: -x[1])
    top20 = ranked[:20]
    rest = ranked[20:]
    for i, (name, amt) in enumerate(top20):
        r = phdr + 1 + i
        ws.cell(row=r, column=1, value=name).font = S.f_body()
        c = ws.cell(row=r, column=2, value=amt)
        c.number_format = '£#,##0.00'
        c.border = S.BORDER_THIN
    if rest:
        r = phdr + 1 + len(top20)
        ws.cell(row=r, column=1, value=f"All other practices ({len(rest)})").font = S.f_body()
        c = ws.cell(row=r, column=2, value=round(sum(a for _, a in rest), 2))
        c.number_format = '£#,##0.00'
        c.border = S.BORDER_THIN
    prac_last = phdr + len(top20) + (1 if rest else 0)

    bar = BarChart()
    bar.type = "col"
    bar.title = "Expenses by practice"
    bar.height, bar.width = 8, 14
    bdata = Reference(ws, min_col=2, min_row=phdr, max_row=prac_last)
    bcats = Reference(ws, min_col=1, min_row=phdr + 1, max_row=prac_last)
    bar.add_data(bdata, titles_from_data=True)
    bar.set_categories(bcats)
    bar.legend = None
    ws.add_chart(bar, f"D{prac_row}")

    # ---------------- By staff type (Category) ----------------
    cat_row = prac_last + 3
    ws[f"A{cat_row}"] = "By staff type"
    ws[f"A{cat_row}"].font = S.f_subheader()
    chdr = cat_row + 1
    ws.cell(row=chdr, column=1, value="Category")
    ws.cell(row=chdr, column=2, value="Total Expenses £")
    S.style_header_row(ws, chdr, 1, 2)
    cat_ranked = sorted(by_category.items(), key=lambda x: -x[1])
    for i, (name, amt) in enumerate(cat_ranked):
        r = chdr + 1 + i
        ws.cell(row=r, column=1, value=name).font = S.f_body()
        c = ws.cell(row=r, column=2, value=amt)
        c.number_format = '£#,##0.00'
        c.border = S.BORDER_THIN
    cat_last = chdr + max(len(cat_ranked), 1)

    donut = DoughnutChart()
    donut.title = "Expenses by staff type"
    donut.height, donut.width = 8, 10
    ddata = Reference(ws, min_col=2, min_row=chdr, max_row=cat_last)
    dcats = Reference(ws, min_col=1, min_row=chdr + 1, max_row=cat_last)
    donut.add_data(ddata, titles_from_data=True)
    donut.set_categories(dcats)
    donut.dataLabels = DataLabelList()
    donut.dataLabels.showPercent = True
    ws.add_chart(donut, f"D{cat_row}")

    # ---------------- Top positions claiming ----------------
    pos_row = cat_last + 20
    ws[f"A{pos_row}"] = "Top positions claiming expenses"
    ws[f"A{pos_row}"].font = S.f_subheader()
    phdr2 = pos_row + 1
    for i, h in enumerate(["Position Name", "Claimant Count", "Total Expenses £"]):
        ws.cell(row=phdr2, column=1 + i, value=h)
    S.style_header_row(ws, phdr2, 1, 3)
    pos_ranked = sorted(by_position.items(), key=lambda x: -x[1][0])[:10]
    for i, (name, (amt, cnt)) in enumerate(pos_ranked):
        r = phdr2 + 1 + i
        ws.cell(row=r, column=1, value=name).font = S.f_body()
        ws.cell(row=r, column=2, value=cnt).border = S.BORDER_THIN
        c = ws.cell(row=r, column=3, value=amt)
        c.number_format = '£#,##0.00'
        c.border = S.BORDER_THIN

    S.col_widths(ws, {"A": 26, "B": 16, "C": 16})
    ws.freeze_panes = "A3"
    return ws


def build_output3(dataset, result, path):
    wb = Workbook()
    wb.remove(wb.active)
    last_row, data_start = build_all_employees_sheet(wb, result)
    build_dashboard_sheet(wb, result, last_row, data_start)
    build_expenses_sheet(wb, dataset, result)
    order = ["Dashboard", "All Employees", "Expenses"]
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    wb.save(path)
    return path


if __name__ == "__main__":
    from mock_data import build_dataset
    from reference_calc import compute_reconciliation
    ds = build_dataset()
    result = compute_reconciliation(ds)
    build_output3(ds, result, "output/Hakim Group - SLT Pack.xlsx")
    print("Output 3 written.")
